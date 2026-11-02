import csv
import io
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.analytics.recompute import recompute
from app.core.auth import require_permission
from app.core.config import settings
from app.core.config_files import app_config, catalog, indicators
from app.core.database import get_session
from app.datasets.ingest import load_file, register_upload, validate_file
from app.datasets.profiles import profiles
from app.models import DatasetFile, IngestionRun
from app.schemas.observatory import DatasetFileItem, DatasetProfileItem, DatasetUploadResult

router = APIRouter(prefix="/datasets", tags=["datasets"])
_pagination = app_config().pagination
ALLOWED_SUFFIXES = (".xlsx", ".csv")
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class CommitRequest(BaseModel):
    # the file repeats values of official sources and the person accepts loading them apart
    confirm_overlap: bool = False


def _profile(code: str):
    profile = profiles().get(code)
    if profile is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND,
                            app_config().messages["unknown_profile"].format(code=code))
    return profile


@router.get("/profiles", response_model=list[DatasetProfileItem])
def list_profiles():
    """Files the platform knows how to read: what to upload and which format is expected."""
    return [DatasetProfileItem(
        dataset_code=p.dataset_code, title=p.title, source=p.source, landing_page=p.landing_page,
        automatic=p.download is not None, requires_declared_source=p.requires_declared_source,
        file_format=p.file.format, sheet=p.file.sheet, expected_headers=p.expected.headers,
        template_columns=p.long_format.columns if p.long_format else None,
        required_columns=[p.long_format.columns[k] for k in p.long_format.required]
        if p.long_format else None,
        mapped_columns={k: v.model_dump(exclude_defaults=True) for k, v in p.columns.items()}
        or None,
        has_template=p.long_format is not None,
        example_files=list(app_config().datasets.example_files.get(p.dataset_code, {})),
    ) for p in profiles().values()]


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, name in enumerate(header):
        width = max([len(str(name)), *(len(str(r[i])) for r in rows if i < len(r) and r[i] is not None)])
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = min(width + 2, 70)
    return ws


@router.get("/profiles/{code}/template", response_class=Response)
def template(code: str, format: Literal["xlsx", "csv"] = "xlsx"):
    """Empty template for a long-format profile (e.g. plantilla_generica): the data sheet with the
    headers only, instructions with example rows, and the valid codes (indicators, countries,
    sectors, occupations) from the catalog. csv: the headers only."""
    prof = _profile(code)
    conf = app_config().datasets
    if not prof.long_format:
        raise HTTPException(status.HTTP_404_NOT_FOUND, conf.messages["no_template"].format(code=code))
    t = conf.template
    headers = list(prof.long_format.columns.values())
    name = t.file_name.format(code=code)
    if format == "csv":
        out = io.StringIO()
        csv.writer(out, delimiter=prof.file.csv_delimiter).writerow(headers)
        return Response(out.getvalue().encode(prof.file.csv_encoding), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{name}.csv"'})
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, t.sheets["data"], headers, [])
    help_ws = _sheet(wb, t.sheets["help"], [t.instructions[0]], [[line] for line in t.instructions[1:]])
    help_ws.append([])
    help_ws.append([t.examples_title])
    help_ws.cell(help_ws.max_row, 1).font = Font(bold=True)
    help_ws.append(headers)
    for ex in t.examples:
        help_ws.append([ex.get(k) for k in prof.long_format.columns])
    h = t.headers
    _sheet(wb, t.sheets["indicators"], [h["code"], h["name"], h["unit"]],
           [[i.code, i.name, i.unit] for i in indicators().indicators])
    _sheet(wb, t.sheets["countries"], [h["code"], h["name"]], [[c.iso3, c.name_es] for c in catalog().countries])
    _sheet(wb, t.sheets["sectors"], [h["code"], h["name"]], [[x.sector_code, x.name] for x in catalog().sectors])
    _sheet(wb, t.sheets["occupations"], [h["code"], h["name"]],
           [[o.occupation_code, o.name] for o in catalog().occupations])
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type=XLSX_TYPE,
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


@router.get("/profiles/{code}/examples/{name}", response_class=FileResponse)
def example_file(code: str, name: str):
    """A file to try the check with (app.yaml -> datasets.example_files)."""
    _profile(code)
    conf = app_config().datasets
    relative = conf.example_files.get(code, {}).get(name)
    if relative is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND,
                            conf.messages["unknown_example"].format(name=name, code=code))
    path = settings.config_dir / relative
    return FileResponse(path, filename=f"{code}_{name}{path.suffix}")


@router.get("", response_model=list[DatasetFileItem])
def list_files(
    profile: str | None = None,
    origin: str | None = Query(None, description="download | upload"),
    limit: int = Query(_pagination.default_limit, ge=1, le=_pagination.max_limit),
    session: Session = Depends(get_session),
):
    """History of every official file read (downloaded or uploaded) and what happened to it."""
    stmt = select(DatasetFile).order_by(DatasetFile.created_at.desc()).limit(limit)
    if profile:
        stmt = stmt.where(DatasetFile.profile_code == profile)
    if origin:
        stmt = stmt.where(DatasetFile.origin == origin)
    return session.scalars(stmt).all()


@router.get("/{dataset_id}", response_model=DatasetFileItem)
def get_file(dataset_id: int, session: Session = Depends(get_session)):
    dataset = session.get(DatasetFile, dataset_id)
    if dataset is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND,
                            app_config().messages["unknown_dataset"].format(id=dataset_id))
    return dataset


@router.post("/upload", response_model=DatasetUploadResult,
             dependencies=[Depends(require_permission("manage_data"))])
async def upload(
    file: UploadFile = File(...),
    profile: str = Form(..., description="dataset_code de /datasets/profiles"),
    declared_source_name: str | None = Form(None),
    declared_source_url: str | None = Form(None),
    reference_date: date | None = Form(None),
    session: Session = Depends(get_session),
):
    """Step 1: store the file and check it. Nothing is loaded until /{id}/commit."""
    conf = app_config().datasets
    prof = _profile(profile)
    name = Path(file.filename or "archivo").name
    if Path(name).suffix.lower() not in ALLOWED_SUFFIXES:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY,
                            conf.messages["unsupported_format"].format(name=name))
    content = await file.read(conf.max_upload_bytes + 1)
    if len(content) > conf.max_upload_bytes:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, conf.messages[
            "too_large"].format(max_mb=round(conf.max_upload_bytes / 1e6)))
    dataset, report = register_upload(session, prof, content, name, declared_source_name,
                                      declared_source_url, reference_date)
    session.commit()
    preview = [r.model_dump() for r in report.records[:conf.preview_rows]] if report.ok else []
    return DatasetUploadResult(**DatasetFileItem.model_validate(dataset).model_dump(),
                               preview=preview)


@router.post("/{dataset_id}/commit", response_model=DatasetFileItem,
             dependencies=[Depends(require_permission("manage_data"))])
def commit(dataset_id: int, body: CommitRequest | None = None, session: Session = Depends(get_session)):
    """Step 2: load a validated upload. A previous load of the same profile (and declared source)
    is replaced. Trends, projections, the index and alerts are recalculated. If the file repeats
    values of official sources, send {"confirm_overlap": true}: they are kept as a separate source."""
    dataset = get_file(dataset_id, session)
    messages = app_config().messages
    if dataset.status != "validated":
        raise HTTPException(status.HTTP_409_CONFLICT,
                            messages["dataset_not_committable"].format(status=dataset.status))
    if (dataset.report.get("overlaps") or {}).get("count") and not (body and body.confirm_overlap):
        raise HTTPException(status.HTTP_409_CONFLICT,
                            app_config().datasets.messages["overlap_confirm_required"])
    prof = _profile(dataset.profile_code)
    report = validate_file(session, prof, Path(dataset.stored_path))  # data may have changed since
    if not report.ok:
        dataset.status, dataset.report = "rejected", report.summary()
        session.commit()
        raise HTTPException(status.HTTP_409_CONFLICT,
                            messages["dataset_not_committable"].format(status="rejected"))
    run = IngestionRun(source_id=dataset.source_id, started_at=datetime.now(UTC),
                       status="running", triggered_by="upload")
    session.add(run)
    session.flush()
    loaded = load_file(session, dataset, prof, report, run)
    run.records_fetched = run.records_normalized = loaded
    session.commit()
    recompute(session, sorted({r.indicator_code for r in report.records}))
    run.status, run.finished_at = "success", datetime.now(UTC)
    session.commit()
    return dataset
