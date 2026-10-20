"""Yearbook-style official files: nested report sheets, one file per year on its own page."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.datasets import ingest
from app.datasets.profiles import profiles
from app.datasets.reader import read_dataset
from app.ingestion.pipeline import run_source
from app.models import DataRevision, DatasetFile, Graduate
from tests.test_datasets import engine  # noqa: F401  (fixture)

PROFILE = profiles()["mec_uy_egresados_grado"]
LIST_URL = PROFILE.landing_page
BOLD = Font(b=True)


def yearbook(year: int, careers: dict[str, list[tuple]], title_year: int | None = None,
             total_extra: float = 0) -> bytes:
    """One sheet per profile sheet: title, header, bold TOTAL, bold headings and plain careers.
    careers: {sheet: [(heading or None, career, graduates)]}; graduates go in column H."""
    wb = Workbook()
    wb.remove(wb.active)
    for spec in PROFILE.outline.sheets:
        ws = wb.create_sheet(f"{spec.sheet}  ")  # real sheet names carry trailing spaces
        ws["A1"] = "OFERTA DE GRADO"
        ws["A5"] = f"EGRESADOS EN PROGRAMAS DE GRADO (Todo el país, {title_year or year})"
        ws["A6"], ws["H6"] = "CARRERA", "TOTAL"
        rows = careers.get(spec.sheet, [])
        total = sum(g for _, _, g in rows if isinstance(g, int | float))
        total += total_extra if spec is PROFILE.outline.sheets[0] else 0
        ws.append([])  # row 7
        ws.append(["", None, None, None, None, None, None, "V"])  # row 8
        ws.append(["TOTAL", 1, 1, 1, 1, 1, 1, total])
        ws["A9"].font = BOLD
        heading = None
        for head, career, graduates in rows:
            if head != heading:
                ws.append([head, 0, 0, 0, 0, 0, 0, 999])  # a heading's own total is not counted
                ws.cell(ws.max_row, 1).font = BOLD
                heading = head
            ws.append([career, 0, 0, 0, 0, 0, 0, graduates])
        ws.append(["Fuente: Oficinas Productoras de cada una de las Instituciones."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CAREERS = {
    "II-7.2.1": [("FACULTAD DE DERECHO", "Abogacía", 200), ("FACULTAD DE DERECHO", "Notariado", 50),
                 ("Salto", "Abogacía ", 10)],
    "II-7.2.2": [("UCU", "Contador Público", 80), ("UCU", "Licenciatura en Cine", "s/d")],
    "II-7.2.3": [("INST UNIV", "Contador  público", 5)],
}


def test_outline_reader_adds_items_by_name_and_checks_the_total(tmp_path):
    path = tmp_path / "a.xlsx"
    path.write_bytes(yearbook(2024, CAREERS))
    r = read_dataset(path, PROFILE)
    assert r.ok and r.periods == ["2024"] and not r.warnings
    got = {(x["career"].strip(), x["level"]): x["value"] for x in r.rows}
    # same program across campuses and institutions (case and spaces aside) is added up;
    # headings and "s/d" are not items
    assert got == {("Abogacía", "Grado"): 210, ("Notariado", "Grado"): 50,
                   ("Contador Público", "Grado"): 85}
    assert all(x["year"] == 2024 for x in r.rows)


def test_outline_reader_small_gap_warns_large_gap_blocks(tmp_path):
    small, large = tmp_path / "s.xlsx", tmp_path / "l.xlsx"
    small.write_bytes(yearbook(2024, CAREERS, total_extra=1))  # 1 of 261 in the first sheet
    large.write_bytes(yearbook(2024, CAREERS, total_extra=100))
    rs, rl = read_dataset(small, PROFILE), read_dataset(large, PROFILE)
    assert rs.ok and len(rs.warnings) == 1
    assert not rl.ok and "suman" in rl.errors[0].message


class YearbookSite:
    """The yearbook list links one page per year; each page links its spreadsheet (relative)."""

    def __init__(self, files: dict[str, tuple[str, bytes]]):
        self.files, self.downloads = files, []  # {year: (file name, content)}

    def __call__(self, url, max_bytes, headers=None):
        base = "/ministerio-educacion-cultura/datos-y-estadisticas/datos/anuario-estadistico-educacion-"
        if url == LIST_URL:
            links = "".join(f'<a href="{base}{y}">Anuario {y}</a>' for y in ["2019", *self.files])
            return links.encode()
        for year, (name, content) in self.files.items():
            if url.endswith(f"anuario-estadistico-educacion-{year}"):
                return (f'<a href="/ministerio-educacion-cultura/sites/ministerio-educacion-cultura/'
                        f'files/2025-04/{name}" download>xlsx</a>').encode()
            if url.endswith(name):
                self.downloads.append(year)
                return content
        raise AssertionError(f"unexpected URL {url}")  # e.g. 2019, before min_release


def _run(engine, site, monkeypatch):  # noqa: F811
    monkeypatch.setattr(ingest, "download_sync", site)
    with Session(engine, expire_on_commit=False) as s:
        return run_source(s, "mec_uruguay", "scheduler")


def _graduates(engine):  # noqa: F811
    with Session(engine) as s:
        return {(g.year, g.career_key): g.graduates for g in s.scalars(select(Graduate)).all()}


def test_one_file_per_year_is_loaded_and_replaced_by_year(engine, monkeypatch):  # noqa: F811
    site = YearbookSite({"2023": ("Anuario_2023.xlsx", yearbook(2023, CAREERS)),
                         "2024": ("Anuario_2024.xlsx", yearbook(2024, CAREERS))})
    run = _run(engine, site, monkeypatch)
    assert run.status == "success" and sorted(site.downloads) == ["2023", "2024"]
    assert _graduates(engine)[(2023, "abogacia")] == 210
    assert {y for y, _ in _graduates(engine)} == {2023, 2024}  # 2019 is before min_release

    _run(engine, site, monkeypatch)  # same links: nothing downloaded
    assert sorted(site.downloads) == ["2023", "2024"]

    # 2024 is republished (new link, revised figures): only 2024 changes
    revised = {**CAREERS, "II-7.2.1": [("FACULTAD DE DERECHO", "Abogacía", 300)]}
    site.files["2024"] = ("Anuario_2024_0.xlsx", yearbook(2024, revised))
    assert _run(engine, site, monkeypatch).status == "success"
    grads = _graduates(engine)
    assert grads[(2024, "abogacia")] == 300 and (2024, "notariado") not in grads
    assert grads[(2023, "abogacia")] == 210 and grads[(2023, "notariado")] == 50
    with Session(engine) as s:
        files = {(f.release, f.file_name): f.status for f in s.scalars(select(DatasetFile)).all()}
    assert files[("2024", "Anuario_2024.xlsx")] == "replaced"
    assert files[("2024", "Anuario_2024_0.xlsx")] == "loaded"
    assert files[("2023", "Anuario_2023.xlsx")] == "loaded"
    with Session(engine) as s:  # the 2024 figures that changed or disappeared are kept
        revs = {(r.record_key["career_key"], r.change): (r.old_value, r.new_value)
                for r in s.scalars(select(DataRevision)).all()}
    # Contador Público is unchanged (other sheets): no revision
    assert revs == {("abogacia", "revised"): (210, 300), ("notariado", "removed"): (50, None)}


def test_file_of_another_year_than_its_page_is_not_loaded(engine, monkeypatch):  # noqa: F811
    site = YearbookSite({"2024": ("Anuario_2024.xlsx", yearbook(2024, CAREERS, title_year=2023))})
    run = _run(engine, site, monkeypatch)
    assert run.status == "format_changed" and "2023" in run.error_message
    assert not _graduates(engine)
