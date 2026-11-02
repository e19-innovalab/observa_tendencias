"""Official files: format check, automatic download with change detection, manual upload."""
import io
from datetime import UTC, date, datetime
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core import config as core_config
from app.core.config_files import app_config, catalog, indicators, sources
from app.core.database import get_session
from app.datasets import ingest
from app.datasets.profiles import profiles
from app.datasets.reader import read_dataset
from app.ingestion.pipeline import run_source
from app.main import app
from app.models import (
    Alert,
    Base,
    Country,
    DatasetFile,
    Indicator,
    Observation,
    Sector,
    Source,
)
from tests.conftest import sqlite_tables

SIPA = profiles()["sipa_asalariados_por_rama"]
HEADERS = SIPA.expected.headers


def sipa_workbook(months: int = 210, rename: dict | None = None, scale_old: float = 1.0) -> bytes:
    """A small file shaped like the SIPA annex: dates up to Jan 2015, then 'feb-15' text,
    provisional '*' months, footnotes after the data."""
    wb = Workbook()
    ws = wb.active
    ws.title = SIPA.file.sheet
    ws.append(["A.2.1. Personas con empleo asalariado registrado ... En miles 1/.", "Volver"])
    ws.append([(rename or {}).get(h, h.upper() if h == "Pesca" else h) for h in HEADERS])
    abbr = {v: k for k, v in SIPA.period.month_abbr.items()}
    for i in range(months):
        year, month = 2009 + i // 12, i % 12 + 1
        if (year, month) <= (2015, 1):
            period = datetime(year, month, 1)
        else:
            period = f"{abbr[month]}-{year % 100:02d}" + ("*" if i >= months - 3 else "")
        factor = scale_old if i < months - 24 else 1.0
        ws.append([period] + [round((100 + i + j) * factor, 3) for j in range(len(HEADERS) - 1)])
    ws.append([])
    ws.append(["* Los datos tienen carácter provisorio"])
    ws.append(["Fuente: MTEySS, SSPyEL, Observatorio de Empleo y Dinámica Empresarial"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_reader_handles_mixed_periods_provisional_and_footnotes(tmp_path):
    path = tmp_path / "sipa.xlsx"
    path.write_bytes(sipa_workbook())
    r = read_dataset(path, SIPA)
    assert r.ok, r.errors
    assert r.periods[0] == "2009-01" and r.periods[-1] == "2026-06" and len(r.periods) == 210
    assert r.provisional_periods == ["2026-04", "2026-05", "2026-06"]
    assert len(r.records) == 210 * len(SIPA.columns)
    assert {x.sector_code for x in r.records} == {"tourism", "health", "energy", None}


def test_reader_blocks_renamed_column(tmp_path):
    path = tmp_path / "sipa.xlsx"
    path.write_bytes(sipa_workbook(rename={"Hoteles y restaurantes": "Alojamiento y gastronomía"}))
    r = read_dataset(path, SIPA)
    assert not r.ok and not r.records
    messages = " ".join(e.message for e in r.errors)
    assert "Las columnas cambiaron" in messages and "alojamiento y gastronomia" in messages


def test_reader_blocks_missing_sheet_and_short_series(tmp_path):
    path = tmp_path / "sipa.xlsx"
    path.write_bytes(sipa_workbook(months=50))
    assert any("Solo hay 50 períodos" in e.message for e in read_dataset(path, SIPA).errors)


@pytest.fixture
def engine(tmp_path, monkeypatch):
    eng = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})
    Base.metadata.create_all(eng, tables=sqlite_tables())
    with Session(eng) as s:
        s.add_all([Country(iso3=c.iso3, name_es=c.name_es, name_en=c.name_en)
                   for c in catalog().countries])
        s.add_all([Source(source_code=x.source_code, name=x.name, source_kind=x.kind,
                          expected_frequency=x.expected_frequency) for x in sources().sources])
        s.add_all([Indicator(indicator_code=i.code, name=i.name, unit=i.unit, domain=i.domain)
                   for i in indicators().indicators])
        s.commit()
    conf = core_config  # files go to a temporary folder
    from app.core.config_files import app_config
    monkeypatch.setattr(app_config().datasets, "raw_dir", tmp_path / "raw")
    monkeypatch.setattr(app_config().datasets, "uploads_dir", tmp_path / "uploads")
    assert conf
    return eng


class FakeSite:
    """Official page that links one file; tracks what was downloaded."""

    def __init__(self, release: str, content: bytes):
        self.release, self.content, self.downloads = release, content, []

    def __call__(self, url, max_bytes, headers=None):
        if url == SIPA.landing_page:
            return (f'<a href="https://www.argentina.gob.ar/sites/default/files/'
                    f'trabajoregistrado_{self.release}_estadisticas.xlsx">Anexo</a>').encode()
        self.downloads.append(url)
        return self.content


def _run(engine, site, monkeypatch):
    monkeypatch.setattr(ingest, "download_sync", site)
    with Session(engine, expire_on_commit=False) as s:
        return run_source(s, "sipa_argentina", "scheduler")


def test_automatic_load_then_skip_same_release(engine, monkeypatch):
    site = FakeSite("2606", sipa_workbook())
    run = _run(engine, site, monkeypatch)
    assert run.status == "success" and run.records_normalized == 210 * 4
    assert len(site.downloads) == 1
    run = _run(engine, site, monkeypatch)  # same link on the page: not even downloaded
    assert run.status == "success" and run.records_normalized == 0 and len(site.downloads) == 1
    with Session(engine) as s:
        f = s.scalar(select(DatasetFile))
        assert f.status == "loaded" and f.release == "2606" and Path(f.stored_path).exists()
        assert f.report["last_period"] == "2026-06" and f.report["fingerprint"]


def test_new_release_with_changed_format_is_not_loaded(engine, monkeypatch):
    _run(engine, FakeSite("2606", sipa_workbook()), monkeypatch)
    with Session(engine) as s:
        before = s.scalar(select(func.count()).select_from(Observation))
    bad = FakeSite("2607", sipa_workbook(months=211, rename={"Total": "Total país"}))
    run = _run(engine, bad, monkeypatch)
    assert run.status == "format_changed" and "columnas cambiaron" in run.error_message
    with Session(engine) as s:
        assert s.scalar(select(func.count()).select_from(Observation)) == before  # untouched
        files = s.scalars(select(DatasetFile).order_by(DatasetFile.dataset_id)).all()
        assert [f.status for f in files] == ["loaded", "format_changed"]
        alert = s.scalar(select(Alert).where(Alert.alert_type == "source_format_changed"))
        assert alert.severity == "high" and "No se cargó" in alert.message


def test_large_revision_of_old_values_blocks_the_load(engine, monkeypatch):
    _run(engine, FakeSite("2606", sipa_workbook()), monkeypatch)
    run = _run(engine, FakeSite("2607", sipa_workbook(months=211, scale_old=1.5)), monkeypatch)
    assert run.status == "format_changed" and "cambiaron más de 5.0%" in run.error_message


def test_new_release_replaces_previous_file(engine, monkeypatch):
    _run(engine, FakeSite("2606", sipa_workbook()), monkeypatch)
    run = _run(engine, FakeSite("2607", sipa_workbook(months=211)), monkeypatch)
    assert run.status == "success"
    with Session(engine) as s:
        files = s.scalars(select(DatasetFile).order_by(DatasetFile.dataset_id)).all()
        assert [f.status for f in files] == ["replaced", "loaded"]
        assert files[1].replaces_dataset_id == files[0].dataset_id


@pytest.fixture
def client(engine, monkeypatch):
    def override():
        with Session(engine) as s:
            yield s

    monkeypatch.setattr(core_config.settings, "admin_api_key", "test-key")
    app.dependency_overrides[get_session] = override
    yield TestClient(app)
    app.dependency_overrides.clear()


KEY = {"X-API-Key": "test-key"}
GOOD_CSV = "indicador,pais,periodo,frecuencia,valor\ninflation,URY,2025,A,4.9\ninflation,CHL,2025,A,4.1\n"


def _upload(client, text, **form):
    data = {"profile": "plantilla_generica", "declared_source_name": "INE Uruguay",
            "declared_source_url": "https://www.gub.uy/instituto-nacional-estadistica", **form}
    return client.post("/api/v1/datasets/upload", headers=KEY, data=data,
                       files={"file": ("datos.csv", text.encode(), "text/csv")})


def test_upload_requires_key_and_reports_row_errors(client):
    assert client.post("/api/v1/datasets/upload", data={"profile": "plantilla_generica"},
                       files={"file": ("x.csv", b"a", "text/csv")}).status_code == 401
    bad = _upload(client, "indicador,pais,periodo,frecuencia,valor\n"
                          "inflation,BRA,2025,A,1\nnope,ARG,2025,A,1\n").json()
    assert bad["status"] == "rejected"
    assert [(e["row"], e["column"]) for e in bad["report"]["errors"]] == [(2, "pais"), (3, "indicador")]
    no_source = _upload(client, GOOD_CSV, declared_source_name="").json()
    assert no_source["status"] == "rejected" and "declarar la fuente" in no_source["report"]["errors"][0]["message"]


def test_upload_preview_commit_and_replace(client, engine):
    up = _upload(client, GOOD_CSV).json()
    assert up["status"] == "validated" and len(up["preview"]) == 2
    with Session(engine) as s:
        assert s.scalar(select(func.count()).select_from(Observation)) == 0  # not loaded yet
    done = client.post(f"/api/v1/datasets/{up['dataset_id']}/commit", headers=KEY).json()
    assert done["status"] == "loaded" and done["records"] == 2
    assert client.post(f"/api/v1/datasets/{up['dataset_id']}/commit", headers=KEY).status_code == 409
    # a new version from the same declared source replaces the previous one's values
    v2 = _upload(client, "indicador,pais,periodo,frecuencia,valor\ninflation,URY,2025,A,5.2\n").json()
    client.post(f"/api/v1/datasets/{v2['dataset_id']}/commit", headers=KEY)
    with Session(engine) as s:
        values = s.scalars(select(Observation.value).where(Observation.data_type == "observed")).all()
        assert values == [5.2]
        statuses = [f.status for f in s.scalars(select(DatasetFile).order_by(DatasetFile.dataset_id))]
        assert statuses == ["replaced", "loaded"]
    history = client.get("/api/v1/datasets", params={"origin": "upload"}).json()
    assert history[0]["declared_source_name"] == "INE Uruguay"


def test_profiles_describe_expected_files(client):
    items = {p["dataset_code"]: p for p in client.get("/api/v1/datasets/profiles").json()}
    sipa = items["sipa_asalariados_por_rama"]
    assert sipa["automatic"] and sipa["sheet"] == "A.2.1" and "Total" in sipa["expected_headers"]
    assert items["plantilla_generica"]["required_columns"] == [
        "indicador", "pais", "periodo", "frecuencia", "valor"]


def test_scheduler_jobs_follow_config():
    from app.jobs.scheduler import build_scheduler, scheduled_sources

    jobs = dict(scheduled_sources())
    assert {"world_bank", "ilostat", "datos_argentina", "sipa_argentina"} <= set(jobs)
    assert "ilo_genai_exposure" not in jobs  # static table: no schedule
    assert jobs["sipa_argentina"] == sources().scheduler.cron_by_frequency["monthly"]
    ids = {j.id for j in build_scheduler().get_jobs()}
    assert "stale_sources" in ids and "source:sipa_argentina" in ids
    assert date and UTC and Sector  # imported for fixtures


def test_cpi_profile_skips_blank_rows_and_empty_cells(tmp_path):
    prof = profiles()["ine_uy_ipc"]
    wb = Workbook()
    ws = wb.active
    ws.title = prof.file.sheet
    for _ in range(prof.file.header_row - 1):
        ws.append(["Índice de Precios del Consumo"])
    ws.append(prof.expected.headers)
    ws.append([])  # blank rows between headers and data, as in the INE file
    ws.append([])
    idx = prof.expected.headers.index("Acum.12 meses")
    for i in range(1100):
        year, month = 1937 + (i + 6) // 12, (i + 6) % 12 + 1
        row = [datetime(year, month, 1)] + [None] * (len(prof.expected.headers) - 1)
        row[1] = 100 + i
        row[idx] = None if i < 12 else 4.5  # no 12-month change for the first year
        ws.append(row)
    ws.append([])
    ws.append(["Fuente: Instituto Nacional de Estadística (INE)."])
    path = tmp_path / "ipc.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(buf.getvalue())
    r = read_dataset(path, prof)
    assert r.ok, r.errors
    assert r.periods[0] == "1937-07" and len(r.periods) == 1100
    assert len(r.records) == 1100 - 12 and {x.country_iso3 for x in r.records} == {"URY"}


def test_files_updated_in_place_are_downloaded_and_compared(engine, monkeypatch):
    prof = profiles()["ine_uy_ipc"]
    assert prof.download.always_download

    class Site:
        downloads = 0

        def __call__(self, url, max_bytes, headers=None):
            if url == prof.landing_page:
                return (b'<a href="https://www5.ine.gub.uy/documents/Estad%C3%ADsticasecon%C3%B3micas/'
                        b'SERIES%20Y%20OTROS/IPC/Base%20Octubre%202022=100/'
                        b'IPC%20gral%20y%20variaciones_base%202022.xlsx">IPC</a>')
            Site.downloads += 1
            return b"not really an xlsx"

    monkeypatch.setattr(ingest, "download_sync", Site())
    for _ in range(2):
        with Session(engine, expire_on_commit=False) as s:
            run_source(s, "ine_uruguay", "scheduler")
    assert Site.downloads == 2  # same URL, but downloaded again to compare


def test_extra_intermediate_certificates_are_loaded():
    from app.connectors.http import ssl_context

    ctx = ssl_context()
    subjects = {dict(x[0] for x in c["subject"]).get("commonName") for c in ctx.get_ca_certs()}
    assert "Sectigo Public Server Authentication CA DV R36" in subjects
    assert ctx.verify_mode.name == "CERT_REQUIRED"


def _diu_workbook(year: int, label: str, disciplines: dict[str, int], total: int) -> bytes:
    """A yearbook workbook: one field-of-study table and the institutions' totals table."""
    import io

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "C 2.1.14"
    ws.append([])
    ws.append([None, f"Cuadro 2.1.14 - Estudiantes, nuevos inscriptos, reinscriptos y egresados de "
                     f"títulos de pregrado y grado por disciplina según institución. Año {year}"])
    ws.append([])
    names, kinds, totals = [None, "Instituciones"], [None, None], [None, "Total general"]
    for name, grads in {"Total Rama": sum(disciplines.values()), **disciplines}.items():
        names += [name, None, None, None]
        kinds += ["EST", "NI", "RE", label]
        totals += [grads * 10, grads * 3, grads * 7, grads]
    ws.append(names)
    ws.append(kinds)
    ws.append(totals)
    t = wb.create_sheet("C 2.1.4")
    t.append([None, f"Cuadro 2.1.4 - Egresados de títulos de pregrado y grado y tasa de crecimiento "
                    f"promedio anual según institución. Período 2014 - {year}"])
    t.append([None, "Instituciones", "Años"])
    t.append([None, None, year - 1, year])
    t.append([None, "Total Instituciones Estatales", total - 5, total])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_diu_yearbook_discipline_tables(tmp_path):
    import zipfile
    from pathlib import Path

    from app.datasets.profiles import profiles
    from app.datasets.reader import read_dataset

    path = Path(tmp_path) / "anuario.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("Z 2.1 Pregrado y Grado Gestión Estatal 24.xlsx",
                   _diu_workbook(2024, "EGRE", {"Informática": 300, "Ingeniería": 200}, 500))
        z.writestr("Z 2.2 Pregrado y Grado Gestión Privada 24.xlsx",
                   _diu_workbook(2024, "EGR", {"Informática": 100}, 100))  # 2023 wording
        z.writestr("Z 4. RHUN 24.xlsx", b"not a workbook we read")
    report = read_dataset(path, profiles()["arg_anuario_diu_egresados"])
    assert report.ok and report.periods == ["2024"]
    rows = {r["career"]: r["value"] for r in report.rows}
    assert rows == {"Informática": 400, "Ingeniería": 200}  # subtotal column left out
    assert {r["level"] for r in report.rows} == {"Pregrado y grado"}
    # the disciplines must add up to the yearbook's own total
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("Z 2.1 Pregrado y Grado Gestión Estatal 24.xlsx",
                   _diu_workbook(2024, "EGRE", {"Informática": 300}, 900))
    assert not read_dataset(path, profiles()["arg_anuario_diu_egresados"]).ok


def test_template_has_headers_codes_and_instructions(client):
    from openpyxl import load_workbook

    r = client.get("/api/v1/datasets/profiles/plantilla_generica/template")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    sheets = app_config().datasets.template.sheets
    assert wb.sheetnames == list(sheets.values())
    data = wb[sheets["data"]]
    assert [c.value for c in data[1]] == ["indicador", "pais", "periodo", "frecuencia", "valor", "sector", "ocupacion"]
    assert data.max_row == 1  # headers only: examples live in the instructions sheet
    assert wb[sheets["indicators"]].max_row - 1 == len(indicators().indicators)
    assert wb[sheets["occupations"]].max_row - 1 == len(catalog().occupations)
    csv_text = client.get("/api/v1/datasets/profiles/plantilla_generica/template", params={"format": "csv"}).text
    assert csv_text.strip().lstrip("﻿") == "indicador,pais,periodo,frecuencia,valor,sector,ocupacion"
    assert client.get("/api/v1/datasets/profiles/sipa_asalariados_por_rama/template").status_code == 404


def test_example_files_pass_and_fail_with_grouped_errors(client):
    good = client.get("/api/v1/datasets/profiles/plantilla_generica/examples/valido").content.decode()
    assert _upload(client, good).json()["status"] == "validated"
    bad = client.get("/api/v1/datasets/profiles/plantilla_generica/examples/con_errores").content.decode()
    report = _upload(client, bad).json()["report"]
    groups = {g["kind"]: g for g in report["groups"]}
    assert groups["unknown_country"]["count"] == 2 and groups["unknown_country"]["values"] == ["BRA", "PER"]
    assert {"unknown_indicator", "bad_period", "bad_frequency", "not_numeric", "missing_value",
            "unknown_occupation"} <= set(groups)
    assert groups["unknown_country"]["hint"] == app_config().datasets.group_hints["unknown_country"]
    assert client.get("/api/v1/datasets/profiles/plantilla_generica/examples/otro").status_code == 404


def test_groups_count_every_error_even_beyond_the_list_limit(client, monkeypatch):
    monkeypatch.setattr(app_config().datasets, "max_errors_reported", 1)
    rows = "".join(f"inflation,BRA,{2000 + i},A,1\n" for i in range(5))
    report = _upload(client, "indicador,pais,periodo,frecuencia,valor\n" + rows).json()["report"]
    assert len(report["errors"]) == 1 and report["groups"][0]["count"] == 5


def _upload_bytes(client, name, content):
    data = {"profile": "plantilla_generica", "declared_source_name": "INE Uruguay",
            "declared_source_url": "https://www.gub.uy/instituto-nacional-estadistica"}
    return client.post("/api/v1/datasets/upload", headers=KEY, data=data,
                       files={"file": (name, content, "application/octet-stream")}).json()


def test_files_that_are_not_what_they_say_are_rejected_before_reading(client, monkeypatch):
    import zipfile

    msgs = app_config().datasets.messages
    wb = Workbook()
    wb.active.append(["indicador", "pais", "periodo", "frecuencia", "valor"])
    real = io.BytesIO()
    wb.save(real)
    with_macros = io.BytesIO(real.getvalue())
    with zipfile.ZipFile(with_macros, "a") as z:
        z.writestr("xl/vbaProject.bin", b"macro")
    cases = [
        ("datos.xlsx", b"esto no es un excel", msgs["not_xlsx"]),
        ("datos.xlsx", with_macros.getvalue(), msgs["has_macros"]),
        ("datos.csv", real.getvalue(), msgs["csv_is_excel"]),
        ("datos.csv", b"indicador\x00\x01\x02", msgs["csv_binary"]),
        ("datos.csv", "indicador,pais\nañ,ARG\n".encode("latin-1"), msgs["csv_encoding"].format(encoding="UTF-8")),
    ]
    for name, content, message in cases:
        up = _upload_bytes(client, name, content)
        assert up["status"] == "rejected" and up["report"]["errors"][0]["message"] == message, name
    monkeypatch.setattr(app_config().datasets, "max_unzipped_bytes", 10)
    assert _upload_bytes(client, "datos.xlsx", real.getvalue())["report"]["errors"][0]["message"].startswith(
        msgs["unzipped_too_large"].split("(")[0])
    monkeypatch.setattr(app_config().datasets, "max_rows", 1)
    assert "más de 1 filas" in _upload(client, GOOD_CSV).json()["report"]["errors"][0]["message"]


def test_values_already_official_need_confirmation_and_stay_apart(client, engine):
    with Session(engine) as s:
        ilostat = s.scalar(select(Source.source_id).where(Source.source_code == "ilostat"))
        s.add(Observation(indicator_code="inflation", country_iso3="URY", period_start=date(2025, 1, 1),
                          period_end=date(2025, 12, 31), frequency="A", value=4.8, unit="percent",
                          data_type="observed", source_id=ilostat, fetched_at=datetime.now(UTC)))
        s.commit()
    up = _upload(client, GOOD_CSV).json()
    overlaps = up["report"]["overlaps"]
    assert up["status"] == "validated" and overlaps["count"] == 1 and overlaps["by_source"] == {"ILOSTAT": 1}
    assert "URY 2025: oficial 4.8, archivo 4.9" in up["report"]["warnings"][0]["message"]
    url = f"/api/v1/datasets/{up['dataset_id']}/commit"
    r = client.post(url, headers=KEY, json={})
    assert r.status_code == 409 and r.json()["detail"] == app_config().datasets.messages["overlap_confirm_required"]
    assert client.post(url, headers=KEY, json={"confirm_overlap": True}).json()["status"] == "loaded"
    with Session(engine) as s:  # the official value is still there, next to the uploaded one
        values = s.execute(select(Source.source_code, Observation.value).join(Source).where(
            Observation.indicator_code == "inflation", Observation.country_iso3 == "URY")).all()
        assert sorted(values) == [("ilostat", 4.8), ("manual_upload", 4.9)]
